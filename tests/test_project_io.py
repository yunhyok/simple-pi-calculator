"""Named project files, migrations and export (DESIGN.md §4.7, §4.8, §5.8.4, §5.8.5, §8.10)."""

from __future__ import annotations

import copy
import json
import os
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from simple_pi_calculator.core.types import DecapRow, LayerRow, PwrRow
from simple_pi_calculator.errors import IssueCollector
from simple_pi_calculator.io import migrations
from simple_pi_calculator.io.export import export_csv, export_xlsx, sheet_name_for
from simple_pi_calculator.io.project_io import (
    Project,
    ProjectFormatError,
    ProjectTooNewError,
    ensure_project_suffix,
    load_project,
    migration_backup_path,
    project_from_dict,
    project_to_dict,
    save_project,
)

TOP_KEY_ORDER = ["format", "schema_version", "app_version", "stackup", "vias", "advanced", "pwr",
                 "decaps", "sweep", "display"]


def _codes(issues) -> list[str]:
    return [i.code for i in issues]


def _sample_project(folder: Path) -> Project:
    p = Project()
    p.stackup_source_path = str(folder / "xl" / "stackup.xlsx")
    p.layers = [LayerRow(1, "TOP", 0.035, 5.8e7, None, None), LayerRow(2, "PP", 0.1, None, 4.2, 0.02),
                LayerRow(3, "GND", 0.035, 5.8e7, 4.3, 0.018)]
    p.vias.drill_diameter_mm = 0.25
    p.vias.pad_via_count = 4
    p.advanced.via_model = "coax"
    p.advanced.mounting_inductance_nh = 0.3
    p.advanced.model_search_dir = str(folder / "models")
    p.pwr_source_path = str(folder / "xl" / "pwr.xlsx")
    p.pwr_rows = [PwrRow("VDD", 3, 1, 40.0, True), PwrRow("VIO", 1, 3, 20.5, False)]
    p.decap_source_path = str(folder / "xl" / "decaps.xlsx")
    p.decap_rows = [
        DecapRow("VDD", str(folder / "models" / "a.mod"), 10, 8.0, dummy=True),
        DecapRow("VIO", str(folder / "b.s2p"), 1, 3.5, dummy=False, subckt="X", s2p_mode="shunt",
                 enabled=False),
    ]
    p.sweep.n_points = 123
    p.sweep.show_plane_only = True
    p.display.z_unit = "uohm"
    return p


# ---------------------------------------------------------------------------------------------
# Round trip and paths
# ---------------------------------------------------------------------------------------------
def test_round_trip_and_relative_paths(tmp_path: Path):
    project = _sample_project(tmp_path)
    path = tmp_path / "board.spical.json"
    save_project(project, path)

    doc = json.loads(path.read_text(encoding="utf-8"))
    assert list(doc.keys()) == TOP_KEY_ORDER
    assert "session" not in doc
    assert doc["stackup"]["source_path"] == "xl/stackup.xlsx"
    assert doc["advanced"]["model_search_dir"] == "models"
    assert doc["decaps"]["rows"][0]["model_file"] == "models/a.mod"
    assert doc["decaps"]["rows"][0]["dummy"] is True
    assert doc["decaps"]["rows"][1]["model_file"] == "b.s2p"
    assert path.read_text(encoding="utf-8").startswith('{\n  "format"')
    assert not list(tmp_path.glob("*.tmp"))

    loaded, issues = load_project(path)
    assert loaded == project
    assert not issues


def test_relative_paths_recomputed_for_new_folder(tmp_path: Path):
    project = _sample_project(tmp_path)
    sub = tmp_path / "deeper" / "dir"
    sub.mkdir(parents=True)
    save_project(project, sub / "copy.spical.json")
    doc = json.loads((sub / "copy.spical.json").read_text(encoding="utf-8"))
    assert doc["stackup"]["source_path"] == "../../xl/stackup.xlsx"
    loaded, _ = load_project(sub / "copy.spical.json")
    assert loaded == project


def test_autosave_style_dict_has_absolute_paths(tmp_path: Path):
    from simple_pi_calculator.io.project_io import Session

    project = _sample_project(tmp_path)
    doc = project_to_dict(project, None, Session(project_path=str(tmp_path / "x.spical.json")))
    assert doc["stackup"]["source_path"] == str(tmp_path / "xl" / "stackup.xlsx").replace("\\", "/")
    assert "session" in doc and list(doc.keys())[-1] == "session"


def test_example_project_loads(example_project_path: Path):
    project, issues = load_project(example_project_path)
    assert not [i for i in issues if i.code.startswith("E_")]
    assert len(project.layers) == 11
    assert [(r.name, r.pwr_layer, r.gnd_layer, r.width_mm) for r in project.pwr_rows] == [
        ("VDD_CORE", 5, 3, 60.0), ("VDD_IO", 7, 9, 30.0)]
    assert [(r.count, r.distance_mm, r.dummy) for r in project.decap_rows] == [
        (10, 8.0, False), (4, 15.0, False), (4, 5.0, True), (1, 10.0, False)]
    assert os.path.isabs(project.decap_rows[0].model_file)
    assert project.vias.drill_diameter_mm == 0.2 and project.vias.via_pitch_mm == 1.0
    assert project.advanced.via_model == "pair"
    assert project.sweep.f_start_hz == 1e5 and project.sweep.n_points == 400
    assert project.display.z_unit == "mohm"
    doc = json.loads(example_project_path.read_text(encoding="utf-8"))
    assert doc["stackup"]["source_path"] == "stackup_6L.xlsx"
    assert doc["decaps"]["rows"][2]["model_file"] == "cap_0402_100nF.mod"


def test_session_in_named_file_ignored_with_info(tmp_path: Path):
    project = _sample_project(tmp_path)
    from simple_pi_calculator.io.project_io import Session

    doc = project_to_dict(project, str(tmp_path), Session(modified=True))
    path = tmp_path / "s.spical.json"
    path.write_text(json.dumps(doc), encoding="utf-8")
    loaded, issues = load_project(path)
    assert loaded == project
    infos = [i for i in issues if i.code == "I_PROJECT_SESSION_IGNORED"]
    assert infos and infos[0].severity.name == "INFO"


def test_unknown_keys_warn(tmp_path: Path):
    doc = project_to_dict(_sample_project(tmp_path), str(tmp_path), None)
    doc["colour"] = "blue"
    doc["vias"]["magic"] = 1
    doc["pwr"]["rows"][0]["height_mm"] = 20
    path = tmp_path / "u.spical.json"
    path.write_text(json.dumps(doc), encoding="utf-8")
    _, issues = load_project(path)
    warned = [i.location for i in issues if i.code == "W_PROJECT_UNKNOWN_KEY"]
    assert set(warned) == {"colour", "vias.magic", "pwr.rows[0].height_mm"}


def test_defaults_for_missing_optional_keys(tmp_path: Path):
    doc = {
        "format": "simple-pi-calculator-project",
        "schema_version": 1,
        "pwr": {"rows": [{"name": "VDD", "pwr_layer": 5, "gnd_layer": 3, "width_mm": 60}]},
        "decaps": {"rows": [{"pwr_name": "VDD", "model_file": "a.mod", "count": 2,
                             "distance_mm": 5}]},
    }
    path = tmp_path / "d.spical.json"
    path.write_text(json.dumps(doc), encoding="utf-8")
    project, issues = load_project(path)
    assert not issues
    row = project.decap_rows[0]
    assert row.dummy is False and row.enabled is True and row.subckt is None and row.s2p_mode is None
    assert row.model_file == str(tmp_path / "a.mod")
    assert project.pwr_rows[0].enabled is True
    assert project.vias.vias_per_decap == 2 and project.vias.antipad_diameter_mm == 0.5
    assert project.advanced.s2p_default_mode == "series"
    assert project.sweep.f_stop_hz == 1e9 and project.sweep.show_plane_only is False
    assert project.layers == []


def test_model_file_resolved_relative_to_excel_folder(tmp_path: Path):
    (tmp_path / "xl").mkdir()
    (tmp_path / "xl" / "cap.mod").write_text("* x\n", encoding="utf-8")
    doc = {"format": "simple-pi-calculator-project", "schema_version": 1,
           "decaps": {"source_path": "xl/decaps.xlsx",
                      "rows": [{"pwr_name": "V", "model_file": "cap.mod", "count": 1,
                                "distance_mm": 1}]}}
    project, _ = project_from_dict(doc, str(tmp_path), IssueCollector())
    assert project.decap_rows[0].model_file == str(tmp_path / "xl" / "cap.mod")


# ---------------------------------------------------------------------------------------------
# Errors
# ---------------------------------------------------------------------------------------------
def test_newer_schema_rejected(tmp_path: Path):
    doc = project_to_dict(Project(), str(tmp_path), None)
    doc["schema_version"] = 2
    path = tmp_path / "n.spical.json"
    path.write_text(json.dumps(doc), encoding="utf-8")
    with pytest.raises(ProjectTooNewError) as exc:
        load_project(path)
    assert exc.value.schema_version == 2
    assert exc.value.issue.code == "E_PROJECT_NEWER"


@pytest.mark.parametrize("content", [
    "{not json",
    json.dumps({"format": "something-else", "schema_version": 1}),
    json.dumps([1, 2, 3]),
    json.dumps({"format": "simple-pi-calculator-project"}),
    json.dumps({"format": "simple-pi-calculator-project", "schema_version": 1,
                "pwr": {"rows": "nope"}}),
    json.dumps({"format": "simple-pi-calculator-project", "schema_version": 1,
                "pwr": {"rows": [{"name": "V", "pwr_layer": "five", "gnd_layer": 3,
                                  "width_mm": 1}]}}),
])
def test_format_errors(tmp_path: Path, content: str):
    path = tmp_path / "bad.spical.json"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ProjectFormatError) as exc:
        load_project(path)
    assert exc.value.issue.code == "E_PROJECT_FORMAT"


def test_invalid_bytes_are_format_error(tmp_path: Path):
    path = tmp_path / "bin.spical.json"
    path.write_bytes(b"\xff\xfe\x00garbage")
    with pytest.raises(ProjectFormatError):
        load_project(path)


# ---------------------------------------------------------------------------------------------
# Migration framework (§5.8.4)
# ---------------------------------------------------------------------------------------------
def _rename_display_key(doc: dict) -> dict:
    out = copy.deepcopy(doc)
    out["display"] = {"z_unit": out.pop("view")["unit"]}
    return out


def test_migration_chain(tmp_path: Path, monkeypatch):
    v1 = project_to_dict(_sample_project(tmp_path), str(tmp_path), None)
    v1.pop("display")
    v1["view"] = {"unit": "ohm"}   # the (hypothetical) version-1 spelling
    path = tmp_path / "old.spical.json"
    path.write_text(json.dumps(v1), encoding="utf-8")

    monkeypatch.setattr(migrations, "CURRENT_SCHEMA_VERSION", 2)
    monkeypatch.setattr(migrations, "MIGRATIONS", {1: _rename_display_key})

    snapshot = copy.deepcopy(v1)
    migrated = _rename_display_key(v1)
    assert v1 == snapshot                      # migration function does not mutate its input
    assert migrated["display"] == {"z_unit": "ohm"}

    project, issues = load_project(path)
    assert "I_PROJECT_MIGRATED" in _codes(issues)
    assert project.display.z_unit == "ohm"
    assert project.migrated_from == 1          # GUI marks the project modified

    original_bytes = path.read_bytes()
    save_project(project, path)
    backup = Path(migration_backup_path(str(path), 1))
    assert backup.name == "old.schema1.bak.spical.json"
    assert backup.read_bytes() == original_bytes
    assert json.loads(path.read_text(encoding="utf-8"))["schema_version"] == 2
    assert project.migrated_from is None


def test_migration_missing_step(monkeypatch):
    monkeypatch.setattr(migrations, "CURRENT_SCHEMA_VERSION", 3)
    monkeypatch.setattr(migrations, "MIGRATIONS", {1: lambda d: dict(d)})
    doc = {"format": "simple-pi-calculator-project", "schema_version": 1}
    with pytest.raises(ProjectFormatError):
        migrations.migrate(doc, IssueCollector())


def test_migrate_current_is_identity():
    doc = {"format": "simple-pi-calculator-project", "schema_version": 1}
    issues = IssueCollector()
    assert migrations.migrate(doc, issues) is doc
    assert not issues.issues


# ---------------------------------------------------------------------------------------------
# Save As suffix rule (§5.8.5)
# ---------------------------------------------------------------------------------------------
@pytest.mark.parametrize("name, expected", [
    ("board", "board.spical.json"), ("board.json", "board.spical.json"),
    ("Board.SPICAL.JSON", "Board.SPICAL.JSON"), ("dir/x.spical.json", "dir/x.spical.json")])
def test_save_as_suffix(name, expected):
    assert ensure_project_suffix(name) == expected


# ---------------------------------------------------------------------------------------------
# Export (§4.8)
# ---------------------------------------------------------------------------------------------
def _fake_result(name: str, plane: bool):
    f = np.geomspace(1e5, 1e9, 5)
    z = (1e-3 + 1j * f * 1e-12)
    return SimpleNamespace(name=name, f_hz=f, z_pad=z, z_plane_only=z * 10 if plane else None,
                           marker_f_hz=np.array([1e6, 1e7]), marker_z=np.array([2e-3, 3e-3j]),
                           info={"C_plane": 4.8e-10, "M": 805})


def test_sheet_names():
    used = {"Summary"}
    assert sheet_name_for("VDD[1]:core/*?", used) == "VDD_1__core___"
    long_name = "X" * 40
    assert sheet_name_for(long_name, used) == "X" * 31
    assert sheet_name_for(long_name, used) == "X" * 29 + "~1"


def test_export_csv(tmp_path: Path):
    paths = export_csv([_fake_result("VDD_CORE", True), _fake_result("VDD_IO", False)],
                       str(tmp_path), "board")
    assert [os.path.basename(p) for p in paths] == ["board_VDD_CORE.csv", "board_VDD_IO.csv"]
    lines = Path(paths[0]).read_text(encoding="utf-8").splitlines()
    data = [ln for ln in lines if not ln.startswith("#")]
    assert data[0] == "Frequency (Hz),Re Z (Ohm),Im Z (Ohm),|Z| (Ohm),|Z| plane only (Ohm)"
    assert data[1].split(",")[0] == "1.000000000e+05"
    assert len(data) == 6
    assert any("|Z| @ 1e+06 Hz" in ln for ln in lines)
    header2 = [ln for ln in Path(paths[1]).read_text(encoding="utf-8").splitlines()
               if not ln.startswith("#")][0]
    assert "plane only" not in header2


def test_export_xlsx(tmp_path: Path):
    import openpyxl

    path = tmp_path / "out.xlsx"
    export_xlsx([_fake_result("VDD_CORE", True), _fake_result("VDD_IO", False)], str(path),
                _sample_project(tmp_path))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Summary", "VDD_CORE", "VDD_IO"]
    ws = wb["VDD_CORE"]
    assert [c.value for c in ws[1]] == ["Frequency (Hz)", "Re Z (Ohm)", "Im Z (Ohm)", "|Z| (Ohm)",
                                        "|Z| plane only (Ohm)"]
    assert ws.max_row == 6
    assert ws["A2"].value == pytest.approx(1e5)


def test_to_inputs_example(example_project_path: Path):
    pytest.importorskip("simple_pi_calculator.core.engine")
    from simple_pi_calculator.io.project_io import to_inputs

    project, _ = load_project(example_project_path)
    inputs = to_inputs(project, str(example_project_path))
    assert len(inputs.stackup.layers) == 11
    assert inputs.vias.drill_diameter_m == pytest.approx(0.2e-3)
    assert inputs.vias.via_pitch_m == pytest.approx(1.0e-3)
    assert [p.name for p in inputs.pwrs] == ["VDD_CORE", "VDD_IO"]
    assert inputs.pwrs[0].width_m == pytest.approx(0.06)
    assert inputs.decap_rows[2].dummy is True
    assert inputs.decap_rows[2].distance_m == pytest.approx(5e-3)
    assert inputs.project_dir == str(example_project_path.parent)
