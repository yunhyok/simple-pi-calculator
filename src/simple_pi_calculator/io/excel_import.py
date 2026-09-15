"""Excel import of the stack-up, PWR list and decap list (DESIGN.md §4.1–§4.4, §5.3).

Every reader adds issues to the collector and returns what could be parsed. Fatal problems (file
cannot be opened, header row not found, unknown length unit) are added to ``issues`` and raised as
:class:`~simple_pi_calculator.errors.InputError`.
"""

from __future__ import annotations

import math
import os
from dataclasses import dataclass, field
from typing import Any, Sequence

from simple_pi_calculator.core.stackup import Layer, Stackup
from simple_pi_calculator.core.types import DecapRow, PwrRow
from simple_pi_calculator.core.units import MM, length_unit_factor_mm
from simple_pi_calculator.errors import InputError, IssueCollector
from simple_pi_calculator.io.excel_headers import (
    DECAP_RULES,
    PWR_RULES,
    STACKUP_RULES,
    ColumnRule,
    Predicate,
    cell_ref,
    find_header_row,
    normalize_sheet_name,
    parse_bool_cell,
    parse_mode_cell,
    parse_number_cell,
    pwr_ignored_column,
)

_BAD = object()  #: sentinel for a cell that produced an error
_DIELECTRIC_TEXT = {"-", "n/a", "na", "none", "—", "–"}


@dataclass
class _Table:
    path: str
    sheet: str
    header_row: int
    mapping: dict[str, tuple[int, str | None]]
    rules: dict[str, ColumnRule]
    rows: list[tuple[int, tuple[Any, ...]]]  # (excel row number, values)
    length_factors: dict[str, float] = field(default_factory=dict)
    _formula_rows: list[tuple[Any, ...]] | None = None
    _formula_loaded: bool = False

    # -- cell access ----------------------------------------------------------------------------
    def raw(self, values: tuple[Any, ...], name: str) -> Any:
        if name not in self.mapping:
            return None
        col = self.mapping[name][0]
        value = values[col] if col < len(values) else None
        if isinstance(value, str) and value.strip() == "":
            return None
        return value

    def ref(self, row: int, name: str) -> str:
        return f"{self.sheet}!{cell_ref(row, self.mapping[name][0])}"

    def is_formula_without_value(self, row: int, name: str) -> bool:
        """True if the cell holds a formula whose cached value is missing (§4.1)."""
        if not self._formula_loaded:
            self._formula_loaded = True
            try:
                import openpyxl

                wb = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
                try:
                    ws = wb[self.sheet]
                    self._formula_rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
                finally:
                    wb.close()
            except Exception:  # pragma: no cover - defensive
                self._formula_rows = None
        if self._formula_rows is None or row - 1 >= len(self._formula_rows):
            return False
        cells = self._formula_rows[row - 1]
        col = self.mapping[name][0]
        value = cells[col] if col < len(cells) else None
        return isinstance(value, str) and value.startswith("=")


def _sheet_rows(ws: Any) -> list[tuple[Any, ...]]:
    """Cell values of a worksheet with merged ranges filled from their top-left cell."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    merged = getattr(ws, "merged_cells", None)  # absent on read-only worksheets
    for rng in list(merged.ranges) if merged is not None else []:
        r0, c0 = rng.min_row - 1, rng.min_col - 1
        if r0 >= len(rows) or c0 >= len(rows[r0]):
            continue
        value = rows[r0][c0]
        for r in range(r0, min(rng.max_row, len(rows))):
            row = rows[r]
            if len(row) < rng.max_col:
                row.extend([None] * (rng.max_col - len(row)))
            for c in range(c0, rng.max_col):
                if row[c] is None:
                    row[c] = value
    return [tuple(r) for r in rows]


def _open_table(path: str | os.PathLike[str], keywords: Sequence[str],
                rules: Sequence[ColumnRule], issues: IssueCollector,
                ignored: Predicate | None = None) -> _Table:
    """Open the workbook, choose the sheet, find the header and collect data rows (§4.1)."""
    import openpyxl

    path_str = os.fspath(path)
    try:
        # Not read-only: merged-cell ranges are only available on normal worksheets (a PWR name
        # merged over several decap rows is common in hand-made sheets).
        wb = openpyxl.load_workbook(path_str, read_only=False, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any openpyxl/zip/IO error
        err = issues.error("E_XL_OPEN", f"Cannot open Excel file: {exc}", path_str)
        raise InputError([err]) from exc
    try:
        # §4.1 sheet choice (first keyword sheet, else the active sheet); if that sheet has no
        # recognisable header, the remaining sheets are tried in workbook order.
        candidates: list[str] = []
        for name in wb.sheetnames:
            if any(k in normalize_sheet_name(name) for k in keywords):
                candidates.append(name)
                break
        active = wb.active.title if wb.active is not None else wb.sheetnames[0]
        for name in [active, *wb.sheetnames]:
            if name not in candidates:
                candidates.append(name)
        sheets = {name: _sheet_rows(wb[name]) for name in candidates}
    finally:
        wb.close()

    header = None
    sheet_name = candidates[0]
    all_rows = sheets[sheet_name]
    for name in candidates:
        if find_header_row(sheets[name], rules, IssueCollector(), name, ignored) is not None:
            sheet_name, all_rows = name, sheets[name]
            break
    header = find_header_row(all_rows, rules, issues, sheet_name, ignored)
    if header is None:
        raise InputError([issues.issues[-1]])

    rule_map = {r.field: r for r in rules}
    table = _Table(path=path_str, sheet=sheet_name, header_row=header.row, mapping=header.mapping,
                   rules=rule_map, rows=[])

    # length unit factors
    unit_errors = []
    for name, (col, unit) in header.mapping.items():
        if rule_map[name].kind == "length":
            try:
                table.length_factors[name] = length_unit_factor_mm(unit)
            except KeyError:
                unit_errors.append(issues.error(
                    "E_XL_UNIT", f"Unknown length unit '{unit}' in column '{rule_map[name].display}'"
                    " (allowed: mm, um, mil, m, in).", path_str,
                    f"{sheet_name}!{cell_ref(header.row, col)}"))
    if unit_errors:
        raise InputError(unit_errors)

    required = [r.field for r in rules if r.required]
    for offset, values in enumerate(all_rows[header.row:]):
        row_no = header.row + 1 + offset
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in values):
            continue
        if all(table.raw(values, name) is None for name in required):
            break
        table.rows.append((row_no, values))
    return table


# ---------------------------------------------------------------------------------------------
# Typed cell parsing
# ---------------------------------------------------------------------------------------------
def _formula_check(table: _Table, row: int, name: str, issues: IssueCollector) -> bool:
    if name not in table.mapping:
        return False
    if table.is_formula_without_value(row, name):
        issues.error("E_XL_FORMULA_NO_VALUE",
                     "Formula cell has no cached value — open and save the file in Excel.",
                     table.path, table.ref(row, name))
        return True
    return False


def _cell_number(table: _Table, values: tuple[Any, ...], row: int, name: str,
                 issues: IssueCollector) -> Any:
    """float, ``None`` (empty) or ``_BAD``."""
    raw = table.raw(values, name)
    if raw is None:
        return _BAD if _formula_check(table, row, name, issues) else None
    try:
        number = parse_number_cell(raw)
    except ValueError:
        issues.error("E_XL_NUMBER", f"Cannot read a number from '{raw}'.", table.path,
                     table.ref(row, name))
        return _BAD
    if number is not None and name in table.length_factors:
        number *= table.length_factors[name]
    return number


def _cell_int(table: _Table, values: tuple[Any, ...], row: int, name: str,
              issues: IssueCollector, code: str = "E_XL_NUMBER") -> Any:
    """int, ``None`` (empty) or ``_BAD``."""
    raw = table.raw(values, name)
    if raw is None:
        return _BAD if _formula_check(table, row, name, issues) else None
    try:
        number = parse_number_cell(raw)
    except ValueError:
        number = float("nan")
    if number is None or not math.isfinite(number) or not float(number).is_integer():
        issues.error(code, f"Expected an integer, found '{raw}'.", table.path, table.ref(row, name))
        return _BAD
    return int(number)


def _cell_str(table: _Table, values: tuple[Any, ...], row: int, name: str,
              issues: IssueCollector) -> Any:
    """Stripped str, ``None`` (empty) or ``_BAD``."""
    raw = table.raw(values, name)
    if raw is None:
        return _BAD if _formula_check(table, row, name, issues) else None
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


def _missing(table: _Table, row: int, name: str, issues: IssueCollector,
             code: str = "E_XL_MISSING_VALUE") -> None:
    issues.error(code, f"Required value '{table.rules[name].display}' is empty.", table.path,
                 table.ref(row, name))


# ---------------------------------------------------------------------------------------------
# Stack-up (§4.2)
# ---------------------------------------------------------------------------------------------
def read_stackup(path: str | os.PathLike[str], issues: IssueCollector) -> Stackup:
    """Read and validate a stack-up table (§4.2). Returns an SI :class:`Stackup`."""
    table = _open_table(path, ("stack",), STACKUP_RULES, issues)
    source = table.path
    layers: list[Layer] = []
    thickness_reported: set[int] = set()

    for row, values in table.rows:
        number = _cell_int(table, values, row, "layer_number", issues, code="E_STACK_LAYER_NUM")
        if number is None:
            _missing(table, row, "layer_number", issues, code="E_STACK_LAYER_NUM")
            continue
        if number is _BAD:
            continue
        name = _cell_str(table, values, row, "layer_name", issues)
        name = "" if name in (None, _BAD) else name

        thickness = _cell_number(table, values, row, "thickness", issues)
        if thickness is None:
            issues.error("E_STACK_THICKNESS", f"Layer {number}: thickness missing.", source,
                         table.ref(row, "thickness"))
            thickness_reported.add(number)
            thickness = 0.0
        elif thickness is _BAD:
            thickness_reported.add(number)
            thickness = 0.0
        elif thickness <= 0.0:
            issues.error("E_STACK_THICKNESS", f"Layer {number}: thickness must be > 0.", source,
                         table.ref(row, "thickness"))
            thickness_reported.add(number)

        raw_sigma = table.raw(values, "conductivity")
        conductivity: float | None
        if isinstance(raw_sigma, str) and raw_sigma.strip().casefold() in _DIELECTRIC_TEXT:
            conductivity = None
        else:
            sigma = _cell_number(table, values, row, "conductivity", issues)
            conductivity = None if sigma in (None, _BAD) or sigma == 0.0 else float(sigma)

        dk = _cell_number(table, values, row, "dk", issues)
        df = _cell_number(table, values, row, "df", issues)
        layers.append(Layer(number=number, name=name, thickness_m=float(thickness) * MM,
                            conductivity=conductivity,
                            dk=None if dk is _BAD else dk, df=None if df is _BAD else df))

    stackup = Stackup(tuple(layers))
    scratch = IssueCollector()
    stackup.validate(scratch, source)
    for issue in scratch.issues:
        if issue.code == "E_STACK_THICKNESS" and any(
                issue.message.startswith((f"Layer {n}:", f"Layer {n} (")) for n in thickness_reported):
            continue
        issues.issues.append(issue)
    return stackup


# ---------------------------------------------------------------------------------------------
# PWR list (§4.3)
# ---------------------------------------------------------------------------------------------
def read_pwr_list(path: str | os.PathLike[str], issues: IssueCollector) -> list[PwrRow]:
    """Read the PWR list (§4.3). Returns mm-valued rows (height is derived, never read)."""
    table = _open_table(path, ("pwr", "power"), PWR_RULES, issues, ignored=pwr_ignored_column)
    source = table.path
    rows: list[PwrRow] = []
    seen: dict[str, int] = {}

    for row, values in table.rows:
        name = _cell_str(table, values, row, "pwr_name", issues)
        pwr_layer = _cell_int(table, values, row, "pwr_layer", issues)
        gnd_layer = _cell_int(table, values, row, "gnd_layer", issues)
        width = _cell_number(table, values, row, "width", issues)
        ok = True
        for field_name, value in (("pwr_name", name), ("pwr_layer", pwr_layer),
                                  ("gnd_layer", gnd_layer)):
            if value is None:
                _missing(table, row, field_name, issues)
                ok = False
            elif value is _BAD:
                ok = False
        if width is None:
            issues.error("E_PWR_DIM", "PWR plane width is missing.", source, table.ref(row, "width"))
            ok = False
        elif width is _BAD:
            ok = False
        if not ok:
            continue
        if width <= 0.0:
            issues.error("E_PWR_DIM", f"PWR '{name}': plane width must be > 0.", source,
                         table.ref(row, "width"))
        if name in seen:
            issues.error("E_PWR_NAME_DUP", f"Duplicate PWR name '{name}' (first in row "
                         f"{seen[name]}).", source, table.ref(row, "pwr_name"))
        else:
            seen[name] = row
        if pwr_layer == gnd_layer:
            issues.error("E_PWR_SAME_LAYER", f"PWR '{name}': PWR and GND layer are both "
                         f"{pwr_layer}.", source, table.ref(row, "pwr_layer"))
        rows.append(PwrRow(name=name, pwr_layer=pwr_layer, gnd_layer=gnd_layer,
                           width_mm=float(width), enabled=True))
    return rows


# ---------------------------------------------------------------------------------------------
# Decap list (§4.4)
# ---------------------------------------------------------------------------------------------
def read_decap_list(path: str | os.PathLike[str], issues: IssueCollector) -> list[DecapRow]:
    """Read the decap assignment list (§4.4), including the optional ``Dummy Cap`` column.

    ``model_file`` is made absolute when it exists relative to the Excel file's folder (the
    first relative candidate of §4.4); otherwise the text is kept for later resolution.
    """
    table = _open_table(path, ("decap", "cap"), DECAP_RULES, issues)
    source = table.path
    excel_dir = os.path.dirname(os.path.abspath(source))
    rows: list[DecapRow] = []

    for row, values in table.rows:
        pwr_name = _cell_str(table, values, row, "pwr_name", issues)
        model_file = _cell_str(table, values, row, "model_file", issues)
        count = _cell_int(table, values, row, "count", issues)
        distance = _cell_number(table, values, row, "distance", issues)
        subckt = _cell_str(table, values, row, "subckt", issues)

        ok = True
        for field_name, value in (("pwr_name", pwr_name), ("model_file", model_file),
                                  ("count", count)):
            if value is None:
                _missing(table, row, field_name, issues)
                ok = False
            elif value is _BAD:
                ok = False
        if distance is None:
            issues.error("E_DECAP_DISTANCE", "Distance to PAD is missing.", source,
                         table.ref(row, "distance"))
            ok = False
        elif distance is _BAD:
            ok = False

        s2p_mode = None
        raw_mode = table.raw(values, "s2p_mode")
        if raw_mode is not None:
            try:
                s2p_mode = parse_mode_cell(raw_mode)
            except ValueError:
                issues.error("E_XL_MODE", f"S2P mode '{raw_mode}' is not 'series' or 'shunt'.",
                             source, table.ref(row, "s2p_mode"))
                ok = False

        dummy = False
        if "dummy" in table.mapping:
            raw_dummy = values[table.mapping["dummy"][0]] if table.mapping["dummy"][0] < len(
                values) else None
            try:
                dummy = parse_bool_cell(raw_dummy)
            except ValueError:
                issues.error("E_XL_BOOL", f"Cannot read a yes/no value from '{raw_dummy}'.",
                             source, table.ref(row, "dummy"))
                ok = False
        if not ok:
            continue

        if distance <= 0.0:
            issues.error("E_DECAP_DISTANCE", f"Distance to PAD must be > 0 (found {distance:g} "
                         "mm).", source, table.ref(row, "distance"))
        if count < 1:
            issues.error("E_DECAP_COUNT", f"Number of decaps must be ≥ 1 (found {count}).",
                         source, table.ref(row, "count"))

        model_text = str(model_file)
        if not os.path.isabs(model_text):
            candidate = os.path.normpath(os.path.join(excel_dir, model_text))
            if os.path.isfile(candidate):
                model_text = candidate
        rows.append(DecapRow(pwr_name=pwr_name, model_file=model_text, count=count,
                             distance_mm=float(distance), dummy=dummy,
                             subckt=None if subckt in (None, _BAD) else subckt,
                             s2p_mode=s2p_mode, enabled=True))
    return rows


__all__ = ["read_stackup", "read_pwr_list", "read_decap_list"]
